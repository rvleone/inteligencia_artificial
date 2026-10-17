"""
Script para geração automatizada do Relatório Acadêmico em formato Word (.docx)
baseado no template institucional 'Template_relatorioIA_AF.docx'.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "Template_relatorioIA_AF.docx")
OUTPUT_DOCX = os.path.join(os.path.dirname(__file__), "docs", "Relatorio_AF_Inteligencia_Artificial.docx")
OUTPUTS_DIR = os.path.join(os.path.dirname(__file__), "outputs")

def set_cell_background(cell, fill_color):
    """Define a cor de fundo de uma célula da tabela."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

def add_heading_styled(doc, text, level):
    """Adiciona títulos formatados com estilo profissional."""
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    for run in h.runs:
        run.font.name = 'Calibri'
        if level == 1:
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D) # Azul escuro institucional
        elif level == 2:
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2B, 0x5C, 0x8F)
        else:
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    return h

def add_paragraph_styled(doc, text, bold_prefix=None, space_after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    """Adiciona parágrafo justificado com entrelinhas e espaçamento ideal."""
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.name = 'Calibri'
        r_bold.font.size = Pt(10.5)
        r_bold.font.bold = True
        
    r_text = p.add_run(text)
    r_text.font.name = 'Calibri'
    r_text.font.size = Pt(10.5)
    r_text.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    return p

def add_image_styled(doc, img_name, caption, width_inches=5.8):
    """Insere imagem centralizada com legenda formatada."""
    img_path = os.path.join(OUTPUTS_DIR, img_name)
    if not os.path.exists(img_path):
        print(f"Aviso: imagem não encontrada: {img_path}")
        return
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(3)
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=Inches(width_inches))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(10)
    r_cap = p_cap.add_run(f"Figura: {caption}")
    r_cap.font.name = 'Calibri'
    r_cap.font.size = Pt(9)
    r_cap.font.italic = True
    r_cap.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

def build_docx():
    print(f"Carregando template: {TEMPLATE_PATH}")
    doc = docx.Document(TEMPLATE_PATH)
    
    # Cabeçalho do Artigo / Relatório
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(18)
    p_title.paragraph_format.space_after = Pt(6)
    r_title = p_title.add_run("APLICAÇÃO DE TÉCNICAS DE APRENDIZADO DE MÁQUINA SUPERVISIONADO E NÃO SUPERVISIONADO NO DIAGNÓSTICO DE CÂNCER DE MAMA (WDBC)")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(15)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("Avaliação Formativa (AF) – Semestre 2026-II | Disciplina: Inteligência Artificial\nUniversidade Salgado de Oliveira (UNIVERSO) – Cursos de Sistemas de Informação / Análise e Desenvolvimento de Sistemas")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(10)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    
    # Autores e Docente
    p_author = doc.add_paragraph()
    p_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_author.paragraph_format.space_after = Pt(16)
    r_author = p_author.add_run("Autor(a): [Insira seu Nome Completo]\nOrientadora: Profa. Dsc. Edhelmira L. M.")
    r_author.font.name = 'Calibri'
    r_author.font.size = Pt(10.5)
    r_author.font.bold = True
    
    # Resumo
    add_heading_styled(doc, "RESUMO", level=1)
    add_paragraph_styled(doc, 
        "Este trabalho apresenta um estudo comparativo rigoroso entre algoritmos de aprendizado de máquina supervisionado e não supervisionado aplicados ao diagnóstico de neoplasias mamárias, utilizando a base de dados pública Breast Cancer Wisconsin Diagnostic (WDBC), proveniente do repositório UCI. A base contempla 569 observações clínicas e 30 atributos morfológicos contínuos extraídos de biópsias aspirativas por agulha fina (PAAF). Inicialmente, realizou-se a análise exploratória e a padronização dos dados por Z-score (StandardScaler). No agrupamento não supervisionado, aplicou-se o algoritmo K-Means, avaliado pelo Método do Cotovelo e Coeficiente de Silhueta, corroborado por projeções 2D e 3D via Análise de Componentes Principais (PCA), alcançando Adjusted Rand Index (ARI) de 0,6707 com k=2. Na etapa de classificação supervisionada, foram implementados e comparados cinco algoritmos fundamentais: Árvore de Decisão, K-Nearest Neighbors (KNN), Support Vector Machine (SVM com kernel RBF), Naive Bayes Gaussiano e Redes Neurais Artificiais (Multi-Layer Perceptron - MLP). A validação foi conduzida sob o protocolo de validação cruzada estratificada em 10 dobras (Stratified 10-Fold Cross-Validation), avaliando acurácia, sensibilidade (recall da classe maligna), especificidade, área sob a curva ROC (AUC) e matriz de confusão. O modelo SVM atingiu o melhor equilíbrio global com 97,54% de acurácia, 96,21% de sensibilidade e AUC de 99,56%, enquanto o KNN obteve a maior especificidade (99,44%). Os resultados demonstram a viabilidade clínica dos modelos como sistemas de suporte à decisão médica, priorizando a minimização de falsos negativos.",
        bold_prefix=None)
    add_paragraph_styled(doc, "Palavras-chave: Inteligência Artificial; Aprendizado de Máquina; Classificação; Clusterização; K-Means; Support Vector Machine; Câncer de Mama.", bold_prefix=None, space_after=14)

    # 1. Introdução
    add_heading_styled(doc, "1. INTRODUÇÃO", level=1)
    add_paragraph_styled(doc, 
        "O câncer de mama configura uma das neoplasias malignas de maior incidência e mortalidade na população feminina globalmente. O prognóstico favorável e as taxas de sobrevida a longo prazo estão diretamente correlacionados com o diagnóstico precoce e a acurácia na diferenciação entre lesões benignas e malignas. Tradicionalmente, o procedimento de Punção Aspirativa por Agulha Fina (PAAF) é empregado para coletar material citológico, o qual é subsequentemente submetido à avaliação morfológica por médicos patologistas.")
    add_paragraph_styled(doc, 
        "Com o advento da Inteligência Artificial (IA) e do Aprendizado de Máquina (Machine Learning), sistemas de auxílio ao diagnóstico médico por computador (CAD - Computer-Aided Diagnosis) tornaram-se ferramentas essenciais para reduzir a variabilidade interobservador, mitigar a fadiga diagnóstica e fornecer segundas opiniões objetivas com alto embasamento estatístico.")
    add_paragraph_styled(doc, 
        "Neste contexto, o presente artigo tem como escopo a implementação, validação empírica e análise crítica comparativa de algoritmos de aprendizado supervisionado (classificação) e não supervisionado (clusterização). O objetivo central é mapear as potencialidades e limitações de cada paradigma analítico frente aos desafios diagnósticos oncológicos, priorizando o equilíbrio entre sensibilidade clínica e especificidade estatística.")

    # 2. Metodologia
    add_heading_styled(doc, "2. METODOLOGIA", level=1)
    
    add_heading_styled(doc, "2.1 Seleção e Descrição do Banco de Dados", level=2)
    add_paragraph_styled(doc, 
        "Para a execução experimental, selecionou-se a base de dados pública Breast Cancer Wisconsin (Diagnostic) - WDBC, mantida no repositório oficial da Universidade da Califórnia em Irvine (UCI Machine Learning Repository). O dataset foi compilado pelo Dr. William H. Wolberg, da Universidade de Wisconsin, e seus colaboradores, através da digitalização de lâminas de PAAF e processamento computacional de contornos celulares.")
    add_paragraph_styled(doc, 
        "A base é composta por 569 instâncias clínicas individuais sem qualquer valor faltante (missing values). Cada amostra contém 30 atributos numéricos contínuos calculados a partir de 10 características geométricas fundamentais dos núcleos celulares: raio, textura, perímetro, área, suavidade, compacidade, concavidade, pontos côncavos, simetria e dimensão fractal. Para cada característica fundamental, reportam-se três medidas estatísticas: média (mean), erro padrão (standard error - se) e o pior valor (worst - média dos três maiores valores medidos).")
    add_paragraph_styled(doc, 
        "A variável alvo (target) é estritamente binária: 357 casos benignos (62,74%) e 212 casos malignos (37,26%). Visando total conformidade com a literatura de epidemiologia clínica, o diagnóstico Maligno foi padronizado como a classe positiva (1) e o diagnóstico Benigno como a classe negativa (0). Dessa forma, a Sensibilidade (Recall) quantifica a capacidade do classificador de identificar os tumores malignos, cujo não reconhecimento (falso negativo) acarreta grave risco de vida.")

    add_heading_styled(doc, "2.2 Pré-Processamento dos Dados", level=2)
    add_paragraph_styled(doc, 
        "Os procedimentos de pré-processamento envolveram verificação de consistência e escalonamento dos atributos. Como os atributos possuem ordens de magnitude amplamente discrepantes (ex.: 'mean area' atinge ordens de 2500, enquanto 'mean smoothness' varia entre 0,05 e 0,16), algoritmos sensíveis à distância euclidiana sofreriam forte viés caso as variáveis não fossem equalizadas.")
    add_paragraph_styled(doc, 
        "Para contornar este efeito, aplicou-se a padronização por Z-score (StandardScaler), transformando cada atributo para ter média zero e desvio-padrão unitário através da equação:")
    add_paragraph_styled(doc, "z = (x - μ) / σ", bold_prefix="Equação 1 (Padronização Z-score): ", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph_styled(doc, 
        "Crucialmente, para mitigar o vazamento de dados (data leakage), a padronização foi estritamente calculada sobre as dobras de treino e replicada nas dobras de validação durante cada ciclo da validação cruzada.")

    add_heading_styled(doc, "2.3 Agrupamento Não Supervisionado (Clusterização)", level=2)
    add_paragraph_styled(doc, 
        "A técnica não supervisionada K-Means foi empregada para avaliar se os dados se agrupam naturalmente em conformidade com o diagnóstico biológico, sem o fornecimento prévio dos rótulos durante o treinamento. A determinação do número ideal de agrupamentos foi orientada por dois critérios matemáticos complementares:")
    add_paragraph_styled(doc, 
        "1. Método do Cotovelo (Elbow Method): monitoramento da Inércia (Soma dos Quadrados Intra-Cluster) em função de k de 2 a 10.\n"
        "2. Coeficiente Médio de Silhueta: métrica que quantifica a coesão interna versus a separabilidade externa de cada ponto, variando no intervalo [-1, 1].")
    add_paragraph_styled(doc, 
        "Adicionalmente, aplicou-se a Análise de Componentes Principais (PCA) para redução de dimensionalidade das 30 dimensões para 2 e 3 componentes ortogonais, permitindo a visualização espacial dos agrupamentos e dos centroides.")

    add_heading_styled(doc, "2.4 Algoritmos de Classificação Supervisionada e Hiperparâmetros", level=2)
    add_paragraph_styled(doc, 
        "Foram selecionados cinco algoritmos representativos de diferentes famílias teóricas de aprendizado de máquina, superando a exigência mínima de três técnicas:")
    add_paragraph_styled(doc, 
        "a) Árvore de Decisão (Decision Tree): Modelo baseado em particionamento recursivo do espaço de características por ganho de informação. Utilizou-se o critério de impureza de Gini com profundidade máxima (max_depth=4) e min_samples_split=5 para podar a complexidade da árvore e preservar interpretabilidade médica.\n"
        "b) K-Nearest Neighbors (KNN): Algoritmo baseado em instâncias que classifica o ponto pela maioria de votos de seus vizinhos mais próximos. Adotou-se k=5 e distância euclidiana (métrica de Minkowski com p=2), configurando um compromisso ideal entre sensibilidade local e imunidade a ruídos pontuais.\n"
        "c) Support Vector Machine (SVM): Classificador baseado na maximização da margem de separação entre hiperplanos. Configurou-se o kernel de Base Radial Gaussiana (RBF) com hiperparâmetro de regularização C=1,0 e escala automática de gamma, viabilizando fronteiras de decisão não-lineares em espaços de dimensão infinita.\n"
        "d) Naive Bayes (Gaussian NB): Modelo probabilístico supervisionado baseado no Teorema de Bayes sob a suposição de independência condicional entre os atributos dada a classe, modelando a densidade contínua através de distribuições Gaussianas.\n"
        "e) Redes Neurais Artificiais (Multi-Layer Perceptron - MLP): Rede neural 'feedforward' densa composta por 2 camadas ocultas com 64 e 32 neurônios respectivamente, função de ativação não linear ReLU, regularização L2 (alpha=0,001) e otimizador estocástico Adam.")

    add_heading_styled(doc, "2.5 Estratégia de Validação e Métricas de Avaliação", level=2)
    add_paragraph_styled(doc, 
        "A avaliação de desempenho foi conduzida mediante Validação Cruzada Estratificada em 10 Dobras (Stratified 10-Fold Cross-Validation). A estratificação assegura que cada partição de teste preserve exatamente a proporção original de classes (62,7% benignos e 37,3% malignos).")
    add_paragraph_styled(doc, 
        "As métricas foram calculadas individualmente por dobra e reportadas pela média amostral acompanhada do desvio-padrão (μ ± σ):")
    add_paragraph_styled(doc, "Acurácia = (VP + VN) / (VP + VN + FP + FN)", bold_prefix="• Acurácia Global: ")
    add_paragraph_styled(doc, "Sensibilidade = VP / (VP + FN)", bold_prefix="• Sensibilidade (Recall): ")
    add_paragraph_styled(doc, "Especificidade = VN / (VN + FP)", bold_prefix="• Especificidade (True Negative Rate): ")
    add_paragraph_styled(doc, "Área sob o gráfico traçado entre a Taxa de Verdadeiros Positivos (Sensibilidade) e a Taxa de Falsos Positivos (1 - Especificidade) através da variação contínua dos limiares de decisão.", bold_prefix="• Área sob a Curva ROC (AUC-ROC): ")
    add_paragraph_styled(doc, "Mapeamento bidimensional cruzado de Verdadeiros Positivos (VP), Verdadeiros Negativos (VN), Falsos Positivos (FP) e Falsos Negativos (FN).", bold_prefix="• Matriz de Confusão: ")

    # 3. Resultados
    add_heading_styled(doc, "3. RESULTADOS E ANÁLISES GRÁFICAS", level=1)
    
    add_heading_styled(doc, "3.1 Análise Exploratória e Distribuição Inicial dos Dados", level=2)
    add_paragraph_styled(doc, 
        "A análise preliminar confirmou a integridade da base, com 0 valores ausentes. A distribuição de classes evidenciou 357 casos benignos e 212 casos malignos, conforme ilustrado na Figura 1.")
    add_image_styled(doc, "eda_distribuicao_classes.png", "Distribuição de Frequências das Classes Diagnósticas (Benigno vs Maligno).", 5.5)

    add_paragraph_styled(doc, 
        "O mapa de calor das correlações lineares (Figura 2) revelou forte colinearidade entre raio, perímetro e área (r > 0,98), o que justifica plenamente a aplicação de técnicas de redução de dimensionalidade como o PCA.")
    add_image_styled(doc, "eda_matriz_correlacao.png", "Matriz de Correlação de Pearson entre Atributos Morfológicos Médios.", 5.0)

    add_paragraph_styled(doc, 
        "A inspeção das distribuições de densidade e boxplots (Figuras 3 e 4) demonstra que as células malignas exibem médias substancialmente mais elevadas em raio, perímetro, concavidade e área em comparação com as células benignas, demonstrando alto potencial de discriminação.")
    add_image_styled(doc, "eda_distribuicoes_atributos.png", "Distribuições de Densidade Estimada (KDE) para Atributos Clínicos Chave.", 5.5)
    add_image_styled(doc, "eda_boxplots.png", "Boxplots Comparativos Padronizados por Diagnóstico.", 5.5)

    add_heading_styled(doc, "3.2 Resultados da Clusterização Não Supervisionada", level=2)
    add_paragraph_styled(doc, 
        "A aplicação do Método do Cotovelo e do Coeficiente de Silhueta (Figura 5) indicou unanimemente k=2 como a partição estrutural ideal, com o maior score de silhueta médio observado (0,3450).")
    add_image_styled(doc, "cluster_elbow_silhouette.png", "Método do Cotovelo e Coeficiente de Silhueta para determinação de k ótimo.", 5.5)

    add_paragraph_styled(doc, 
        "A projeção 2D via PCA (Figura 6) capturou 63,2% da variância total acumulada nas duas primeiras componentes principais (PC1=44,3% e PC2=19,0%). A comparação direta entre os agrupamentos encontrados pelo K-Means e o diagnóstico clínico real revelou um Adjusted Rand Index (ARI) de 0,6707, comprovando que os atributos celulares possuem uma separação espontânea marcante.")
    add_image_styled(doc, "cluster_pca_2d.png", "Projeção PCA 2D: Comparativo entre Clusters K-Means e Classes Diagnósticas Reais.", 5.8)

    add_paragraph_styled(doc, 
        "Na projeção 3D (Figura 7), a inclusão da terceira componente principal elevou a variância explicada acumulada para 72,6%, permitindo visualizar a clara segregação espacial entre os perfis sadios e tumorais.")
    add_image_styled(doc, "cluster_pca_3d.png", "Visualização Espacial Tridimensional dos Dados Clínicos via PCA 3D.", 5.0)

    add_heading_styled(doc, "3.3 Desempenho Comparativo dos Modelos de Classificação", level=2)
    add_paragraph_styled(doc, 
        "A Tabela 1 sintetiza as métricas obtidas pelos 5 modelos sob o protocolo de validação cruzada estratificada em 10 dobras (Stratified 10-Fold CV).")

    # Tabela de Resultados
    table = doc.add_table(rows=6, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    headers = ["Algoritmo", "Acurácia (%)", "Sensibilidade (%)", "Especificidade (%)", "AUC-ROC (%)"]
    widths = [Inches(1.8), Inches(1.1), Inches(1.1), Inches(1.1), Inches(1.1)]

    # Cabeçalho
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].width = widths[i]
        set_cell_background(hdr_cells[i], "1B365D")
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = 'Calibri'
            r.font.bold = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Dados da tabela
    data_rows = [
        ["Árvore de Decisão", "91,04 ± 3,87", "85,80 ± 10,24", "94,13 ± 4,56", "91,45 ± 7,34"],
        ["K-Nearest Neighbors (KNN)", "97,01 ± 1,76", "92,90 ± 4,90", "99,44 ± 1,11", "98,68 ± 1,46"],
        ["Support Vector Machine (SVM)", "97,54 ± 1,95", "96,21 ± 3,56", "98,33 ± 2,83", "99,56 ± 0,55"],
        ["Naive Bayes Gaussiano", "93,15 ± 3,08", "89,57 ± 7,65", "95,26 ± 4,82", "98,56 ± 1,12"],
        ["Redes Neurais (MLP)", "96,84 ± 1,89", "95,24 ± 5,63", "97,75 ± 2,45", "99,26 ± 1,00"]
    ]

    for row_idx, row_data in enumerate(data_rows, start=1):
        row_cells = table.rows[row_idx].cells
        bg_col = "F5F8FA" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            row_cells[col_idx].text = text
            row_cells[col_idx].width = widths[col_idx]
            set_cell_background(row_cells[col_idx], bg_col)
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(9)
                if row_idx == 3: # Destacar SVM como melhor
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_tab_note = doc.add_paragraph()
    p_tab_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tab_note.paragraph_format.space_before = Pt(3)
    p_tab_note.paragraph_format.space_after = Pt(12)
    r_tn = p_tab_note.add_run("Tabela 1: Métricas de Desempenho Comparativo (Média ± Desvio Padrão em 10 Folds).")
    r_tn.font.name = 'Calibri'
    r_tn.font.size = Pt(8.5)
    r_tn.font.italic = True

    add_image_styled(doc, "comparativo_metricas_barras.png", "Comparativo Geral das Quatro Métricas Principais por Modelo.", 5.8)

    add_heading_styled(doc, "3.4 Análise das Curvas ROC e Matrizes de Confusão", level=2)
    add_paragraph_styled(doc, 
        "A Figura 9 apresenta a sobreposição das curvas ROC médias obtidas na validação cruzada. O classificador SVM apresentou a maior capacidade de discriminação com AUC de 99,56%, seguido pela Rede Neural MLP com 99,26% e KNN com 98,68%.")
    add_image_styled(doc, "curvas_roc_comparativas.png", "Curvas ROC Comparativas dos Modelos de Classificação Supervisionada.", 5.2)

    add_paragraph_styled(doc, 
        "A Figura 10 expõe as matrizes de confusão agregadas de todos os 10 folds de teste (totalizando as 569 observações). O modelo SVM gerou apenas 8 falsos negativos e 6 falsos positivos em todo o conjunto de dados. Já o KNN obteve expressivo resultado com apenas 2 falsos positivos entre os 357 casos benignos.")
    add_image_styled(doc, "matrizes_confusao.png", "Matrizes de Confusão Agregadas dos Cinco Classificadores Testados.", 5.8)

    # 4. Discussão
    add_heading_styled(doc, "4. DISCUSSÃO CRÍTICA COMPARATIVA", level=1)
    add_paragraph_styled(doc, 
        "A interpretação dos resultados deve ser pautada pelo contexto clínico de aplicação. Em tarefas de diagnóstico oncológico, o custo de um Falso Negativo (paciente com câncer classificada como sadia) é desproporcionalmente mais grave do que o de um Falso Positivo (paciente saudável submetida a exames adicionais). Dessa maneira, a Sensibilidade (Recall) é a métrica mestra de validação.")
    add_paragraph_styled(doc, 
        "Neste quesito, o algoritmo Support Vector Machine (SVM com kernel RBF) demonstrou clara superioridade, registrando 96,21% de sensibilidade e acurácia de 97,54%. A capacidade do kernel RBF de mapear os dados em um espaço de dimensão infinita permitiu isolar fronteiras de decisão complexas e não-lineares, capturando com exatidão os tumores malignos atípicos.")
    add_paragraph_styled(doc, 
        "A Rede Neural Artificial (MLP) apresentou comportamento semelhante, com 95,24% de sensibilidade e 99,26% de AUC. Entretanto, redes neurais possuem maior custo computacional de convergência e operam como caixas-pretas de menor interpretabilidade direta.")
    add_paragraph_styled(doc, 
        "Por sua vez, o KNN destacou-se pela excepcional especificidade de 99,44%, errando apenas 2 casos benignos em 357 pacientes. Isso indica que a vizinhança euclidiana dos casos benignos é extremamente coesa. Contudo, sua sensibilidade de 92,90% implicou em 15 falsos negativos, quase o dobro do SVM (8 falsos negativos).")
    add_paragraph_styled(doc, 
        "O classificador Naive Bayes Gaussiano demonstrou robustez probabilística (AUC de 98,56%), porém sua acurácia de 93,15% reflete a violação da premissa teórica de independência condicional, haja vista as intensas correlações demonstradas na análise exploratória (ex.: r=0,99 entre raio e perímetro). Por fim, a Árvore de Decisão com max_depth=4 apresentou a menor acurácia (91,04%), porém oferece a máxima transparência com regras de decisão explicáveis e diretamente auditáveis.")
    add_paragraph_styled(doc, 
        "No que tange à clusterização, o K-Means alcançou índice ARI de 0,6707 de forma puramente não supervisionada, comprovando que as medições biométricas celulares carregam uma forte assinatura patológica imanente, capaz de guiar a estratificação clínica inicial sem necessidade de prévia anotação médica.")

    # 5. Conclusões
    add_heading_styled(doc, "5. CONCLUSÕES", level=1)
    add_paragraph_styled(doc, 
        "O presente trabalho atingiu plenamente os objetivos propostos na Avaliação Formativa de Inteligência Artificial, explorando o ciclo completo de ciência de dados e machine learning: desde a análise exploratória, passando pela clusterização não supervisionada, até o treinamento, validação cruzada k-fold e avaliação multivariada de 5 modelos de classificação.")
    add_paragraph_styled(doc, 
        "Os achados empíricos consolidam o modelo Support Vector Machine (SVM com kernel RBF) como o mais apto e confiável para aplicação como sistema auxiliar de triagem diagnóstica, minimizando os falsos negativos para apenas 3,77% (sensibilidade de 96,21%) e exibindo AUC de 99,56%.")
    add_paragraph_styled(doc, 
        "Como trabalhos futuros e perspectivas de expansão científica, recomenda-se: (i) a aplicação de técnicas de seleção de atributos baseadas em LASSO (regularização L1) para reduzir redundâncias entre medidas correlatas; (ii) o emprego de comitês de classificadores (Ensemble Learning como Random Forest e XGBoost); e (iii) a otimização fina de hiperparâmetros por Busca Bayesiana.")

    # 6. Referências
    add_heading_styled(doc, "6. REFERÊNCIAS BIBLIOGRÁFICAS", level=1)
    refs = [
        "DUDA, R. O.; HART, P. E.; STORK, D. G. Pattern Classification. 2. ed. New York: John Wiley & Sons, 2001.",
        "FAWCETT, T. An introduction to ROC analysis. Pattern Recognition Letters, v. 27, n. 8, p. 861-874, 2006.",
        "HASTIE, T.; TIBSHIRANI, R.; FRIEDMAN, J. The Elements of Statistical Learning: Data Mining, Inference, and Prediction. 2. ed. New York: Springer, 2009.",
        "PEDREGOSA, F. et al. Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, v. 12, p. 2825-2830, 2011.",
        "ROUSSEEUW, P. J. Silhouettes: a graphical aid to the interpretation and validation of cluster analysis. Journal of Computational and Applied Mathematics, v. 20, p. 53-65, 1987.",
        "STREET, W. N.; WOLBERG, W. H.; MANGASARIAN, O. L. Nuclear feature extraction for breast tumor diagnosis. In: IS&T/SPIE 1993 International Symposium on Electronic Imaging: Science and Technology, v. 1905, p. 861-870, 1993.",
        "WOLBERG, W. H.; STREET, W. N.; MANGASARIAN, O. L. Breast Cancer Wisconsin (Diagnostic) Data Set. UCI Machine Learning Repository, 1995. Disponível em: <https://archive.ics.uci.edu/ml/datasets/Breast+Cancer+Wisconsin+(Diagnostic)>."
    ]
    for ref in refs:
        p_ref = doc.add_paragraph()
        p_ref.paragraph_format.space_after = Pt(4)
        p_ref.paragraph_format.line_spacing = 1.0
        r_ref = p_ref.add_run(ref)
        r_ref.font.name = 'Calibri'
        r_ref.font.size = Pt(9)
        r_ref.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    os.makedirs(os.path.dirname(OUTPUT_DOCX), exist_ok=True)
    doc.save(OUTPUT_DOCX)
    print(f"Relatório DOCX gerado com sucesso em: {OUTPUT_DOCX}")

if __name__ == "__main__":
    build_docx()
